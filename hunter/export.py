"""The order sheet a seller sends to RhinoMall, built from their cart (hunter.orders).

An .xlsx when openpyxl is installed, a UTF-8 CSV (which Excel opens) otherwise, so the site
works either way. The sheet has the seller's details at the top — the same ones RhinoMall's
order needs — then one row per product with the 1688 link, the quantity, the cost to Dubai,
the RhinoMall selling price and the profit, and a totals row.
"""

from __future__ import annotations

import csv
import io
from datetime import datetime, timezone
from typing import Any

# Column headers, in order. The value for each comes from _row below.
COLUMNS = [
    "ردیف",
    "نام محصول",
    "دسته‌بندی",
    "کلیدواژه چینی",
    "کلیدواژه انگلیسی",
    "لینک 1688",
    "لینک مرجع (آمازون/تمو)",
    "قیمت واحد (یوان)",
    "تعداد در لیستینگ",
    "حداقل سفارش",
    "تعداد سفارش",
    "تمام‌شده دبی ($)",
    "قیمت فروش ($)",
    "سود هر واحد ($)",
    "سرمایه‌ی این قلم ($)",
    "درآمد انتظاری ($)",
    "سود انتظاری ($)",
    "حاشیه سود",
    "یادداشت",
]


def _cat_fa(item: dict[str, Any]) -> str:
    from .categories import CATEGORIES

    key = item.get("category") or ""
    return CATEGORIES[key].fa if key in CATEGORIES else key


def _row(i: int, it: dict[str, Any]) -> list[Any]:
    return [
        i,
        it.get("title_fa") or "محصول",
        _cat_fa(it),
        it.get("keyword_zh") or "",
        it.get("keyword_en") or "",
        it.get("product_url") or "",
        it.get("reference_url") or "",
        round(float(it.get("price_cny") or 0), 2),
        int(it.get("units") or 1),
        int(it.get("moq") or 1),
        int(it.get("qty") or 0),
        it.get("landed_usd") or 0,
        it.get("sell_usd") or 0,
        it.get("profit_usd") or 0,
        it.get("order_capital_usd") or 0,
        it.get("order_income_usd") or 0,
        it.get("order_profit_usd") or 0,
        f"{round((it.get('margin') or 0) * 100)}%",
        it.get("note") or "",
    ]


def _header_lines(
    user: dict[str, Any], summary: dict[str, Any], toman_per_usd: float
) -> list[list[Any]]:
    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    toman = int(round(summary.get("capital_usd", 0) * toman_per_usd, -4))
    return [
        ["سفارش شکارچی برای راینومال"],
        ["نام فروشنده", user.get("name") or ""],
        ["کد ملی", user.get("national_id") or ""],
        ["موبایل", user.get("phone") or ""],
        ["ایمیل", user.get("email") or ""],
        ["تاریخ", today],
        ["نرخ دلار (تومان)", int(toman_per_usd)],
        ["تعداد اقلام", summary.get("count", 0)],
        ["سرمایه‌ی کل ($)", summary.get("capital_usd", 0), "تومان", toman],
        [],
    ]


def _totals_row(summary: dict[str, Any]) -> list[Any]:
    row: list[Any] = [""] * len(COLUMNS)
    row[1] = "جمع کل"
    row[10] = summary.get("total_qty", 0)
    row[14] = summary.get("capital_usd", 0)
    row[15] = summary.get("income_usd", 0)
    row[16] = summary.get("profit_usd", 0)
    row[17] = f"{round(summary.get('margin', 0) * 100)}%"
    return row


def build_order(
    user: dict[str, Any], cart: dict[str, Any], toman_per_usd: float
) -> tuple[bytes, str, str]:
    """Return (bytes, filename, media_type). An .xlsx if openpyxl is here, else a CSV."""
    items = cart.get("items", [])
    summary = cart.get("summary", {})
    header = _header_lines(user, summary, toman_per_usd)
    rows = [_row(i, it) for i, it in enumerate(items, 1)]
    try:
        return _xlsx(header, rows, summary), "sefaresh-rhinomall.xlsx", XLSX_MIME
    except ImportError:
        return _csv(header, rows, summary), "sefaresh-rhinomall.csv", "text/csv; charset=utf-8"


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# A cell starting with one of these is run as a formula by Excel or LibreOffice when the sheet
# is a CSV ("CSV injection"); in an .xlsx only "=" is, when the cell is stored as a formula.
# The sheet goes from the seller to RhinoMall's staff, so nothing a seller typed (a product
# name, a note) may ever run on their computers.
FORMULA_START = ("=", "+", "-", "@", "\t", "\r", "\n", "＝", "＋", "－", "＠")


def _csv_safe(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(FORMULA_START):
        return "'" + value
    return value


def _xlsx(header: list[list[Any]], rows: list[list[Any]], summary: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "سفارش"
    ws.sheet_view.rightToLeft = True
    bold = Font(bold=True)
    for line in header:
        ws.append(line)
        if line and line[0]:
            ws.cell(row=ws.max_row, column=1).font = bold
    head_at = ws.max_row + 1
    ws.append(COLUMNS)
    fill = PatternFill("solid", fgColor="1F2937")
    white = Font(bold=True, color="FFFFFF")
    for col in range(1, len(COLUMNS) + 1):
        c = ws.cell(row=head_at, column=col)
        c.fill, c.font = fill, white
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    ws.append(_totals_row(summary))
    for col in range(1, len(COLUMNS) + 1):
        c = ws.cell(row=ws.max_row, column=col)
        c.font = bold
    widths = [6, 26, 16, 18, 20, 30, 30, 14, 14, 12, 12, 14, 13, 14, 16, 16, 16, 11, 24]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=head_at + 1, column=1)
    for line in ws.iter_rows():
        for c in line:
            if isinstance(c.value, str):  # text stays text, never a formula
                c.data_type = "s"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _csv(header: list[list[Any]], rows: list[list[Any]], summary: dict[str, Any]) -> bytes:
    buf = io.StringIO()
    buf.write("﻿")  # BOM, so Excel reads the Persian as UTF-8
    w = csv.writer(buf)
    for line in [*header, COLUMNS, *rows, _totals_row(summary)]:
        w.writerow([_csv_safe(v) for v in line])
    return buf.getvalue().encode("utf-8")
